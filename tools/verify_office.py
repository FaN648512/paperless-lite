# -*- coding: utf-8 -*-
"""
Office 格式（Excel / PPT）导入验收
=================================
1. 现造 4 类测试文件：.xlsx / .xls / .pptx / .ppt（老版二进制）
2. 直接调用处理管线验证文本抽取（不依赖服务）
3. 检查关键内容是否抽到、页数/表数是否正确
"""
import os
import sys

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)

OUT = os.path.join(HERE, "tools", "officetest")
THUMBS = os.path.join(HERE, "tools", "_thumbs")   # 缩略图单独放，避免污染扫描目录
os.makedirs(OUT, exist_ok=True)
os.makedirs(THUMBS, exist_ok=True)

from core.pipeline import run_pipeline  # noqa: E402

PASS = FAIL = 0


def check(name, ok, extra=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print("  [通过] %s %s" % (name, extra))
    else:
        FAIL += 1
        print("  [失败] %s %s" % (name, extra))


# ---------------- 造测试文件 ----------------
def make_xlsx(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "检测项目清单"
    ws.append(["序号", "检测项目", "检测点位", "结果", "单位"])
    ws.append([1, "苯", "喷漆车间", 0.5, "mg/m3"])
    ws.append([2, "甲苯", "喷漆车间", 1.2, "mg/m3"])
    ws.append([3, "噪声", "破碎机房", 86.5, "dB(A)"])
    ws2 = wb.create_sheet("汇总")
    ws2.append(["受检单位", "蓝盾检测技术有限公司"])
    ws2.append(["联系人", "张伟 13800000000"])
    wb.save(path)


def make_xls(path):
    import xlwt
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("培训计划")
    ws.write(0, 0, "2026年度安全生产培训计划")
    ws.write(1, 0, "期数")
    ws.write(1, 1, "主题")
    ws.write(2, 0, 1)
    ws.write(2, 1, "作业场所危害因素识别与防护")
    wb.save(path)


def make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "安全生产培训课件"
    s1.placeholders[1].text = "主讲：蓝盾检测技术有限公司"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "作业场所危害因素识别"
    s2.placeholders[1].text = "化学因素\n物理因素\n生物因素"
    prs.save(path)


def make_ppt_legacy(path):
    """
    造一个"像老版 .ppt"的二进制：正文在真实 .ppt 里以 UTF-16LE
    存放在 TextCharsAtom 记录中，这里用同样的编码塞进垃圾字节之间，
    用来验证尽力提取逻辑能否把文字捞出来。
    """
    junk = bytes(range(256))
    chunks = [b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", junk,
              "劳动合同管理规范".encode("utf-16-le"), junk[:64],
              "甲方：示例市某建材加工厂".encode("utf-16-le"), junk[:32],
              "合同期限三年".encode("utf-16-le"), junk]
    with open(path, "wb") as f:
        f.write(b"".join(chunks))


def main():
    files = {}
    make_xlsx(os.path.join(OUT, "01_检测项目清单.xlsx"))
    make_xls(os.path.join(OUT, "02_培训计划.xls"))
    make_pptx(os.path.join(OUT, "03_安全生产培训.pptx"))
    make_ppt_legacy(os.path.join(OUT, "04_劳动合同管理.ppt"))

    print("=== 1. 直接调用处理管线 ===")
    expect = {
        "01_检测项目清单.xlsx": (["苯", "喷漆车间", "蓝盾检测技术有限公司"], "excel"),
        "02_培训计划.xls": (["作业场所危害因素识别与防护", "2026年度安全生产培训计划"], "excel"),
        "03_安全生产培训.pptx": (["安全生产培训课件", "作业场所危害因素识别", "化学因素"], "ppt"),
        "04_劳动合同管理.ppt": (["劳动合同管理规范", "合同期限三年"], "ppt"),
    }
    for name, (keys, group) in expect.items():
        path = os.path.join(OUT, name)
        ext = name.rsplit(".", 1)[1].lower()
        thumb = os.path.join(THUMBS, "_t_%s.jpg" % ext)
        try:
            content, source, count = run_pipeline(path, ext, thumb)
        except Exception as e:
            check("%s 解析" % name, False, "异常：%s" % e)
            continue
        hit = all(k in content for k in keys)
        check("%s 解析（%s，%d 表/页，%d 字）" % (name, source, count, len(content)),
              hit and source == "office" and count >= 1,
              "" if hit else "缺失关键词：%s" % [k for k in keys if k not in content])
        check("%s 生成缩略图" % name, os.path.isfile(thumb) and os.path.getsize(thumb) > 0)

    print("\n结果：%d 通过 / %d 失败" % (PASS, FAIL))
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
